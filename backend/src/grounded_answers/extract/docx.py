"""
DOCX into the common structure (ADR 0020, decision 4).

The body is read in document order. A heading style (Heading 1 to 9, or Title) starts a section;
its level comes from the numbering when the heading has one, otherwise from the style. Tables are
kept as rows and attached to the section they sit in.

The XML inside a DOCX is parsed by python-docx with resolve_entities=False (checked in the source,
and tested with an XXE and a billion-laughs payload in tests/test_extract_docx.py). Nothing here
changes that.
"""

import io
import re

from docx import Document
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph

from grounded_answers.extract.model import (
    Extracted,
    Section,
    Table,
    level_from_number,
    split_number,
)

_HEADING_STYLE: re.Pattern[str] = re.compile(r"^Heading (\d)$")


def _style_level(style_name: str) -> int | None:
    if style_name == "Title":
        return 1
    match = _HEADING_STYLE.match(style_name)
    if match is None:
        return None
    return min(int(match.group(1)), 6)


class _Builder:
    def __init__(self, number: str | None, heading: str, level: int) -> None:
        self.number = number
        self.heading = heading
        self.level = level
        self.text: list[str] = []
        self.tables: list[Table] = []

    def build(self) -> Section:
        return Section(
            number=self.number,
            heading=self.heading,
            level=self.level,
            page=None,
            text="\n".join(part for part in self.text if part.strip()),
            tables=tuple(self.tables),
        )


def extract_docx(data: bytes) -> Extracted:
    document = Document(io.BytesIO(data))
    sections: list[_Builder] = []
    current = _Builder(None, "", 1)
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            style = paragraph.style.name if paragraph.style is not None else ""
            style_level = _style_level(style)
            if style_level is not None and paragraph.text.strip():
                number, title = split_number(paragraph.text)
                level = level_from_number(number) if number else style_level
                sections.append(current)
                current = _Builder(number, title, level)
            elif paragraph.text.strip():
                current.text.append(paragraph.text.strip())
        elif tag == "tbl":
            table = DocxTable(child, document)
            rows = tuple(tuple(cell.text.strip() for cell in row.cells) for row in table.rows)
            current.tables.append(Table(rows=rows))
    sections.append(current)
    built = tuple(
        section.build() for section in sections if section.heading or section.text or section.tables
    )
    return Extracted(format="docx", page_count=None, sections=built)
