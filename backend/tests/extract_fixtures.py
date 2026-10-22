"""
Generated test inputs for extraction (ADR 0020). Tests never read the demo data.

Everything here is built in code, so the fixtures are reproducible and carry no real documents.
"""

import io
import zlib
from collections.abc import Iterable, Iterator

from docx import Document


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf_pages(pages: Iterable[Iterable[tuple[float, str]]]) -> bytes:
    """
    A PDF with one content stream per page. Each line is (font size, text), placed top to bottom.
    The file has a correct xref table, so every parser reaches the content.
    """
    objects: list[bytes] = [b"", b""]  # 1 catalog, 2 pages (filled in last)
    page_numbers: list[int] = []
    font_number = 0
    for lines in pages:
        content = bytearray()
        for index, (size, text) in enumerate(lines):
            y = 760 - index * (size + 12)
            content += f"BT /F1 {size} Tf 72 {y} Td ({_escape(text)}) Tj ET\n".encode("latin-1")
        content_number = len(objects) + 1
        objects.append(
            f"<< /Length {len(content)} >>\nstream\n".encode() + bytes(content) + b"\nendstream"
        )
        page_number = len(objects) + 1
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {{FONT}} 0 R >> >> /Contents {content_number} 0 R >>".encode()
        )
        page_numbers.append(page_number)
    font_number = len(objects) + 1
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{number} 0 R" for number in page_numbers)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_numbers)} >>".encode()
    objects = [obj.replace(b"{FONT}", str(font_number).encode()) for obj in objects]
    return assemble(objects)


def assemble(objects: list[bytes]) -> bytes:
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    size = len(objects) + 1
    out += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Root 1 0 R /Size {size} >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def blank_pdf() -> bytes:
    """A valid PDF with no text at all: the scanned-PDF case (D-C)."""
    return assemble(
        [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] >>",
        ]
    )


def docx_with_headings() -> bytes:
    """A DOCX with heading styles, numbered and unnumbered, a paragraph and a table."""
    document = Document()
    document.add_heading("Expense Policy", 0)
    document.add_heading("4 Hours of Driving", 1)
    document.add_paragraph("Drivers keep a log of every shift.")
    document.add_heading("4.2 Rest Breaks", 2)
    document.add_paragraph("A break is due after five hours.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Shift"
    table.cell(0, 1).text = "Break"
    table.cell(1, 0).text = "Day"
    table.cell(1, 1).text = "30 minutes"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def stream_bomb(inflated_bytes: int) -> bytes:
    """A zlib stream that inflates to inflated_bytes of zeros. Built in chunks, so the build is small."""
    compressor = zlib.compressobj(9)
    chunk = bytes(1024 * 1024)
    out = bytearray()
    remaining = inflated_bytes
    while remaining > 0:
        size = min(len(chunk), remaining)
        out += compressor.compress(chunk[:size])
        remaining -= size
    out += compressor.flush()
    return bytes(out)


def numbered_markdown(sections: Iterator[tuple[str, str]]) -> str:
    return "\n\n".join(f"{heading}\n\n{body}" for heading, body in sections)
