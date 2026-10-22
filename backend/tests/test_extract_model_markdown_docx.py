"""The common structure, and the Markdown and DOCX extractors (ADR 0020). No database, no subprocess."""

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document

from grounded_answers.extract.docx import extract_docx
from grounded_answers.extract.markdown import extract_markdown, extract_text_file
from grounded_answers.extract.model import (
    Extracted,
    Section,
    level_from_number,
    split_number,
)
from tests.extract_fixtures import docx_with_headings

# ---- the common structure -----------------------------------------------------------------------


@pytest.mark.parametrize(
    ("number", "level"), [("4", 1), ("4.2", 2), ("4.2.3", 3), ("1.2.3.4.5.6.7", 6)]
)
def test_the_level_of_a_numbered_heading_is_its_number_of_parts(number: str, level: int) -> None:
    assert level_from_number(number) == level


@pytest.mark.parametrize(
    ("heading", "number", "title"),
    [
        ("4.2.3 Vehicle checks", "4.2.3", "Vehicle checks"),
        ("4. Hours of driving", "4", "Hours of driving"),
        ("Driver Handbook", None, "Driver Handbook"),
        ("4.2", None, "4.2"),
    ],
)
def test_a_heading_splits_into_number_and_title(
    heading: str, number: str | None, title: str
) -> None:
    assert split_number(heading) == (number, title)


@pytest.mark.parametrize("level", [0, 7])
def test_a_section_level_outside_one_to_six_is_refused(level: int) -> None:
    with pytest.raises(ValueError):
        Section(number=None, heading="x", level=level, page=None, text="")


def test_a_section_number_must_look_like_a_number() -> None:
    with pytest.raises(ValueError):
        Section(number="4.a", heading="x", level=1, page=None, text="")


def test_a_page_number_starts_at_one() -> None:
    with pytest.raises(ValueError):
        Section(number=None, heading="x", level=1, page=0, text="")


def test_a_document_with_only_whitespace_has_no_text() -> None:
    assert not Extracted(
        format="txt", page_count=None, sections=(Section(None, "", 1, None, "  \n"),)
    ).has_text()


# ---- Markdown ----------------------------------------------------------------------------------


MARKDOWN = """# Driver Handbook

Intro paragraph.

## 4 Hours of Driving

Drivers keep a log.

### 4.2 Rest breaks

A break is due after five hours.

| Shift | Break |
|---|---|
| Day | 30 minutes |

```
code line
```
"""


def test_markdown_numbered_headings_take_their_level_from_the_number() -> None:
    result = extract_markdown(MARKDOWN)
    by_number = {section.number: section for section in result.sections if section.number}
    assert by_number["4"].level == 1
    assert by_number["4.2"].level == 2
    assert by_number["4.2"].heading == "Rest breaks"


def test_markdown_unnumbered_title_takes_its_tag_level() -> None:
    result = extract_markdown(MARKDOWN)
    title = next(section for section in result.sections if section.heading == "Driver Handbook")
    assert title.number is None
    assert title.level == 1
    assert "Intro paragraph." in title.text


def test_markdown_table_is_kept_as_rows_on_its_section() -> None:
    result = extract_markdown(MARKDOWN)
    rest = next(section for section in result.sections if section.number == "4.2")
    assert rest.tables[0].rows == (("Shift", "Break"), ("Day", "30 minutes"))


def test_markdown_code_block_is_kept_as_text() -> None:
    result = extract_markdown(MARKDOWN)
    rest = next(section for section in result.sections if section.number == "4.2")
    assert "code line" in rest.text


def test_markdown_text_before_the_first_heading_is_a_section_with_no_number() -> None:
    result = extract_markdown("Preamble text.\n\n# Title\n\nBody.")
    assert result.sections[0].number is None
    assert result.sections[0].heading == ""
    assert result.sections[0].text == "Preamble text."


def test_plain_text_is_one_section_with_no_structure() -> None:
    result = extract_text_file("  Leave policy: 25 days.\n\n4.2 looks like a number but is text.  ")
    assert result.format == "txt"
    assert len(result.sections) == 1
    assert result.sections[0].number is None
    assert result.sections[0].heading == ""
    assert (
        result.sections[0].text == "Leave policy: 25 days.\n\n4.2 looks like a number but is text."
    )


# ---- DOCX --------------------------------------------------------------------------------------


def test_docx_headings_come_from_their_styles_and_numbers() -> None:
    result = extract_docx(docx_with_headings())
    titles = [(section.number, section.heading, section.level) for section in result.sections]
    assert (None, "Expense Policy", 1) in titles
    assert ("4", "Hours of Driving", 1) in titles
    assert ("4.2", "Rest Breaks", 2) in titles


def test_docx_paragraphs_stay_with_their_section_and_tables_are_rows() -> None:
    result = extract_docx(docx_with_headings())
    rest = next(section for section in result.sections if section.number == "4.2")
    assert rest.text == "A break is due after five hours."
    assert rest.tables[0].rows == (("Shift", "Break"), ("Day", "30 minutes"))


def test_docx_with_no_text_has_no_text() -> None:
    buffer = io.BytesIO()
    Document().save(buffer)
    assert not extract_docx(buffer.getvalue()).has_text()


# ---- DOCX XML security (ADR 0020, decision 4) ---------------------------------------------------


XXE_CANARY = "CANARY-7f3a9-do-not-appear"

XXE_BODY = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<!DOCTYPE w:document [<!ENTITY xxe SYSTEM "{path}">]>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body><w:p><w:r><w:t>before &xxe; after</w:t></w:r></w:p></w:body>
</w:document>"""

LAUGHS_BODY = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<!DOCTYPE w:document [
<!ENTITY lol "lol">
<!ENTITY lol1 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">
<!ENTITY lol2 "&lol1;&lol1;&lol1;&lol1;&lol1;&lol1;&lol1;&lol1;&lol1;&lol1;">
<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">
<!ENTITY lol4 "&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;">
<!ENTITY lol5 "&lol4;&lol4;&lol4;&lol4;&lol4;&lol4;&lol4;&lol4;&lol4;&lol4;">
<!ENTITY lol6 "&lol5;&lol5;&lol5;&lol5;&lol5;&lol5;&lol5;&lol5;&lol5;&lol5;">
<!ENTITY lol7 "&lol6;&lol6;&lol6;&lol6;&lol6;&lol6;&lol6;&lol6;&lol6;&lol6;">
<!ENTITY lol8 "&lol7;&lol7;&lol7;&lol7;&lol7;&lol7;&lol7;&lol7;&lol7;&lol7;">
<!ENTITY lol9 "&lol8;&lol8;&lol8;&lol8;&lol8;&lol8;&lol8;&lol8;&lol8;&lol8;">
]>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body><w:p><w:r><w:t>start &lol9; end</w:t></w:r></w:p></w:body>
</w:document>"""


def _docx_with_body(body: str) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        archive.writestr("word/document.xml", body)
    return buffer.getvalue()


def test_an_external_entity_is_not_resolved_so_no_file_content_appears(tmp_path: Path) -> None:
    """The canary file holds a secret. The XXE payload points at it. The secret must not come out."""
    canary = tmp_path / "canary.txt"
    canary.write_text(XXE_CANARY, encoding="utf-8")
    payload = _docx_with_body(XXE_BODY.format(path=canary.as_posix()))
    try:
        result = extract_docx(payload)
    except Exception as exc:  # a refusal is also acceptable: the secret still does not appear
        assert XXE_CANARY not in str(exc)
        return
    text = " ".join(section.text + section.heading for section in result.sections)
    assert XXE_CANARY not in text


def test_the_xxe_check_would_catch_a_leak_when_entities_are_resolved(tmp_path: Path) -> None:
    """Control: a parser that resolves entities does leak the canary, so the check above is sensitive."""
    from lxml import etree  # type: ignore[import-untyped]  # lxml ships no stubs; see ADR 0020 follow-up

    canary = tmp_path / "canary.txt"
    canary.write_text(XXE_CANARY, encoding="utf-8")
    parser = etree.XMLParser(resolve_entities=True)
    root = etree.fromstring(XXE_BODY.format(path=canary.as_posix()).encode("utf-8"), parser)
    assert XXE_CANARY in "".join(root.itertext())


def test_a_billion_laughs_payload_is_refused_or_not_expanded() -> None:
    """Nested entities would expand to about 3 GB. Either a refusal, or a short text: never the blow-up."""
    payload = _docx_with_body(LAUGHS_BODY)
    try:
        result = extract_docx(payload)
    except Exception as exc:
        assert "lol" not in str(exc)
        return
    text = " ".join(section.text for section in result.sections)
    assert len(text) < 1000, "entity expansion must not happen"
